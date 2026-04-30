# #!/usr/bin/env python3
# """
# export_functional_exchanges_to_excel.py
# Exports functional exchanges from a SysML v2 file to an Excel file.
# Extracts:
# - Function IDs from action def doc comments
# - Port names exactly as they appear in port definitions (e.g., FIP1_24b3)
# - Port IDs from port doc comments
# - Exchange IDs from connection doc comments
# """

# import pathlib
# import re
# import sys
# from typing import Dict, List

# import openpyxl
# from openpyxl.styles import PatternFill

# def from_pascal_case(name: str) -> str:
#     """Convert PascalCase to space-separated words with proper handling of acronyms."""
#     if not name:
#         return name
#     if name.isupper():
#         return name

#     result = []
#     i = 0
#     n = len(name)
#     while i < n:
#         if name[i].isupper():
#             j = i
#             while j < n and name[j].isupper():
#                 j += 1
#             if j == n:
#                 result.append(name[i:j])
#                 i = j
#             elif j < n and name[j].islower():
#                 result.append(name[i:j-1])
#                 result.append(" " + name[j-1])
#                 i = j
#             else:
#                 result.append(name[i:j])
#                 i = j
#         else:
#             result.append(name[i])
#             i += 1
#     return "".join(result).replace("  ", " ").strip()

# def extract_id_from_text(text: str) -> str:
#     """Extract ID from text using the pattern: /* ID: uuid */"""
#     match = re.search(r"/\*\s*ID:\s*([0-9a-f-]+)\s*\*/", text)
#     return match.group(1) if match else ""

# def get_block_content(text: str, start_pos: int) -> tuple:
#     """Get content of a block starting with {, handling nested braces."""
#     brace_count = 1
#     end_pos = start_pos
#     while end_pos < len(text) and brace_count > 0:
#         if text[end_pos] == '{':
#             brace_count += 1
#         elif text[end_pos] == '}':
#             brace_count -= 1
#         end_pos += 1
#     return text[start_pos:end_pos-1] if brace_count == 0 else "", end_pos

# def parse_sysml_file(file_path: pathlib.Path) -> List[Dict]:
#     """
#     Parse the SysML file to extract functional exchange definitions.
#     Returns a list of dictionaries with functional exchange information.
#     """
#     with open(file_path, 'r', encoding='utf-8') as f:
#         content = f.read()

#     # First, create a dictionary of all action definitions with their ports and IDs
#     action_info = {}
#     for action_match in re.finditer(r'action def (\w+)\s*\{', content):
#         action_name = action_match.group(1)
#         start = action_match.end()
#         action_content, end = get_block_content(content, start)

#         # Extract action ID
#         action_id = extract_id_from_text(action_content)

#         # Find all port definitions in this action
#         action_info[action_name] = {
#             'id': action_id,
#             'ports': {}
#         }

#         for port_match in re.finditer(r'port (\w+)\s*:\s*\w+\s*\{', action_content):
#             port_usage_name = port_match.group(1)
#             port_start = port_match.end()
#             port_content, port_end = get_block_content(action_content, port_start)

#             # Extract port ID
#             port_id = extract_id_from_text(port_content)

#             action_info[action_name]['ports'][port_usage_name] = {
#                 'id': port_id,
#                 'name': port_usage_name  # Store the exact port usage name
#             }

#     # Find all connection definitions (functional exchanges)
#     functional_exchanges = []
#     for conn_match in re.finditer(r'connection def (\w+)\s*\{', content):
#         conn_name = conn_match.group(1)
#         start = conn_match.end()
#         conn_content, end = get_block_content(content, start)

#         # Extract connection ID
#         conn_id = extract_id_from_text(conn_content)

#         # Extract end actions
#         end_action_matches = re.findall(r'end action (\w+)\s*:\s*(\w+);', conn_content)
#         if len(end_action_matches) < 2:
#             continue

#         from_action_usage, from_action_type = end_action_matches[0]
#         to_action_usage, to_action_type = end_action_matches[1]

#         # Extract interface and connection details
#         interface_match = re.search(
#             r'interface : \w+\s+connect\s+(\w+)\.(\w+)\s+to\s+(\w+)\.(\w+)',
#             conn_content
#         )
#         if not interface_match:
#             continue

#         # The connection is from to_port to from_port (order in the connect statement)
#         to_action_name_in_conn, to_port_usage, from_action_name_in_conn, from_port_usage = interface_match.groups()

#         # Get from action info
#         from_action_data = action_info.get(from_action_type, {})
#         from_action_id = from_action_data.get('id', "")
#         from_port_data = from_action_data.get('ports', {}).get(from_port_usage, {})
#         from_port_id = from_port_data.get('id', "")
#         from_port_name = from_port_data.get('name', "")  # Exact port usage name

#         # Get to action info
#         to_action_data = action_info.get(to_action_type, {})
#         to_action_id = to_action_data.get('id', "")
#         to_port_data = to_action_data.get('ports', {}).get(to_port_usage, {})
#         to_port_id = to_port_data.get('id', "")
#         to_port_name = to_port_data.get('name', "")  # Exact port usage name

#         functional_exchanges.append({
#             "exchange_id": conn_id,
#             "exchange_name": conn_name,
#             "from_action_id": from_action_id,
#             "from_action_name": from_action_type,
#             "from_port_id": from_port_id,
#             "from_port_name": from_port_name,  # Exact port usage name (e.g., FIP1_24b3)
#             "to_action_id": to_action_id,
#             "to_action_name": to_action_type,
#             "to_port_id": to_port_id,
#             "to_port_name": to_port_name  # Exact port usage name (e.g., FOP1_ac6c)
#         })

#     return functional_exchanges

# def create_workbook() -> openpyxl.Workbook:
#     """Create a new workbook with headers."""
#     wb = openpyxl.Workbook()

#     if wb.active.title == "Sheet":
#         wb.remove(wb.active)

#     ws = wb.create_sheet("Functional Exchanges")

#     headers = [
#         "Function From ID", "Function From Name", "Function From Port ID", "Function From Port Name",
#         "Functional Exchange ID", "Functional Exchange Name",
#         "Function To ID", "Function To Name", "Function To Port ID", "Function To Port Name"
#     ]

#     for col, header in enumerate(headers, 1):
#         ws.cell(row=1, column=col, value=header)

#     for cell in ws[1]:
#         cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

#     return wb

# def export_to_excel(functional_exchanges: List[Dict], output_path: pathlib.Path) -> None:
#     """Export functional exchanges to an Excel file."""
#     output_path.parent.mkdir(parents=True, exist_ok=True)
#     wb = create_workbook()
#     ws = wb["Functional Exchanges"]

#     for exchange in functional_exchanges:
#         new_row = ws.max_row + 1

#         # From function and port
#         ws.cell(row=new_row, column=1, value=exchange["from_action_id"])
#         ws.cell(row=new_row, column=2, value=from_pascal_case(exchange["from_action_name"]))
#         ws.cell(row=new_row, column=3, value=exchange["from_port_id"])
#         ws.cell(row=new_row, column=4, value=exchange["from_port_name"])  # Exact port usage name

#         # Functional exchange info
#         ws.cell(row=new_row, column=5, value=exchange["exchange_id"])
#         ws.cell(row=new_row, column=6, value=from_pascal_case(exchange["exchange_name"]))

#         # To function and port
#         ws.cell(row=new_row, column=7, value=exchange["to_action_id"])
#         ws.cell(row=new_row, column=8, value=from_pascal_case(exchange["to_action_name"]))
#         ws.cell(row=new_row, column=9, value=exchange["to_port_id"])
#         ws.cell(row=new_row, column=10, value=exchange["to_port_name"])  # Exact port usage name

#     # Auto-adjust column widths
#     for column in ws.columns:
#         max_length = 0
#         column_letter = column[0].column_letter
#         for cell in column:
#             try:
#                 if len(str(cell.value)) > max_length:
#                     max_length = len(str(cell.value))
#             except:
#                 pass
#         adjusted_width = (max_length + 2) * 1.2
#         ws.column_dimensions[column_letter].width = adjusted_width

#     wb.save(output_path)
#     print(f"Exported {len(functional_exchanges)} functional exchanges to: {output_path}")

# def main() -> None:
#     _DVS_DIR = pathlib.Path(__file__).parent.parent
#     DEFAULT_INPUT = _DVS_DIR / "Functions_generated.sysml"
#     DEFAULT_OUTPUT = _DVS_DIR / "results" / "DVS_Functional_Exchanges_export.xlsx"

#     args = sys.argv[1:]
#     if args and args[0] in ("-h", "--help"):
#         print(__doc__)
#         sys.exit(0)

#     input_path = pathlib.Path(args[0]) if args else DEFAULT_INPUT
#     if not input_path.exists():
#         print(f"Error: file not found: {input_path}", file=sys.stderr)
#         sys.exit(1)

#     output_path = DEFAULT_OUTPUT if len(args) < 2 else pathlib.Path(args[1])

#     print(f"Reading: {input_path}")
#     functional_exchanges = parse_sysml_file(input_path)

#     if not functional_exchanges:
#         print("Warning: No functional exchanges found in the file!")
#     else:
#         print(f"Found {len(functional_exchanges)} functional exchanges in the file")
#         # Print debug info for the first few exchanges
#         for i, exchange in enumerate(functional_exchanges[:3], 1):
#             print(f"\nSample exchange {i}:")
#             print(f"  From: {exchange['from_action_name']}.{exchange['from_port_name']} (ID: {exchange['from_action_id']}, Port ID: {exchange['from_port_id']})")
#             print(f"  To:   {exchange['to_action_name']}.{exchange['to_port_name']} (ID: {exchange['to_action_id']}, Port ID: {exchange['to_port_id']})")
#             print(f"  Exchange: {exchange['exchange_name']} (ID: {exchange['exchange_id']})")

#     export_to_excel(functional_exchanges, output_path)

# if __name__ == "__main__":
#     main()

#!/usr/bin/env python3
"""
export_functional_exchanges_to_excel.py
Exports functional exchanges from a SysML v2 file to an Excel file.
Formats port names for export by:
- Removing the ID suffix (everything after and including the first underscore)
- Adding spaces between letters and numbers (e.g., FOP1 -> FOP 1)
"""

import pathlib
import re
import sys
from typing import Dict, List

import openpyxl
from openpyxl.styles import PatternFill

def from_pascal_case(name: str) -> str:
    """Convert PascalCase to space-separated words with proper handling of acronyms."""
    if not name:
        return name
    if name.isupper():
        return name

    result = []
    i = 0
    n = len(name)
    while i < n:
        if name[i].isupper():
            j = i
            while j < n and name[j].isupper():
                j += 1
            if j == n:
                result.append(name[i:j])
                i = j
            elif j < n and name[j].islower():
                result.append(name[i:j-1])
                result.append(" " + name[j-1])
                i = j
            else:
                result.append(name[i:j])
                i = j
        else:
            result.append(name[i])
            i += 1
    return "".join(result).replace("  ", " ").strip()

def format_port_name_for_export(port_usage_name: str) -> str:
    """Convert port usage name for export by removing ID suffix and formatting.
    Example: 'FOP1_1974' -> 'FOP 1', 'outPublishexperimentresults_55b4' -> 'outPublishexperimentresults'
    """
    # Remove everything after and including the first underscore
    base_name = port_usage_name.split('_')[0]

    # Insert space between letters and numbers
    formatted = re.sub(r'([a-zA-Z])([0-9])', r'\1 \2', base_name)
    formatted = re.sub(r'([0-9])([a-zA-Z])', r'\1 \2', formatted)

    return formatted

def extract_id_from_text(text: str) -> str:
    """Extract ID from text using the pattern: /* ID: uuid */"""
    match = re.search(r"/\*\s*ID:\s*([0-9a-f-]+)\s*\*/", text)
    return match.group(1) if match else ""

def get_block_content(text: str, start_pos: int) -> tuple:
    """Get content of a block starting with {, handling nested braces."""
    brace_count = 1
    end_pos = start_pos
    while end_pos < len(text) and brace_count > 0:
        if text[end_pos] == '{':
            brace_count += 1
        elif text[end_pos] == '}':
            brace_count -= 1
        end_pos += 1
    return text[start_pos:end_pos-1] if brace_count == 0 else "", end_pos

def parse_sysml_file(file_path: pathlib.Path) -> List[Dict]:
    """
    Parse the SysML file to extract functional exchange definitions.
    Returns a list of dictionaries with functional exchange information.
    """
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    # First, create a dictionary of all action definitions with their ports and IDs
    action_info = {}
    for action_match in re.finditer(r'action def (\w+)\s*\{', content):
        action_name = action_match.group(1)
        start = action_match.end()
        action_content, end = get_block_content(content, start)

        # Extract action ID
        action_id = extract_id_from_text(action_content)

        # Find all port definitions in this action
        action_info[action_name] = {
            'id': action_id,
            'ports': {}
        }

        for port_match in re.finditer(r'port (\w+)\s*:\s*\w+\s*\{', action_content):
            port_usage_name = port_match.group(1)
            port_start = port_match.end()
            port_content, port_end = get_block_content(action_content, port_start)

            # Extract port ID
            port_id = extract_id_from_text(port_content)

            action_info[action_name]['ports'][port_usage_name] = {
                'id': port_id,
                'name': port_usage_name  # Store the exact port usage name for ID lookup
            }

    # Find all connection definitions (functional exchanges)
    functional_exchanges = []
    for conn_match in re.finditer(r'connection def (\w+)\s*\{', content):
        conn_name = conn_match.group(1)
        start = conn_match.end()
        conn_content, end = get_block_content(content, start)

        # Extract connection ID
        conn_id = extract_id_from_text(conn_content)

        # Extract end actions
        end_action_matches = re.findall(r'end action (\w+)\s*:\s*(\w+);', conn_content)
        if len(end_action_matches) < 2:
            continue

        from_action_usage, from_action_type = end_action_matches[0]
        to_action_usage, to_action_type = end_action_matches[1]

        # Extract interface and connection details
        interface_match = re.search(
            r'interface : \w+\s+connect\s+(\w+)\.(\w+)\s+to\s+(\w+)\.(\w+)',
            conn_content
        )
        if not interface_match:
            continue

        # The connection is from to_port to from_port (order in the connect statement)
        to_action_name_in_conn, to_port_usage, from_action_name_in_conn, from_port_usage = interface_match.groups()

        # Get from action info
        from_action_data = action_info.get(from_action_type, {})
        from_action_id = from_action_data.get('id', "")
        from_port_data = from_action_data.get('ports', {}).get(from_port_usage, {})
        from_port_id = from_port_data.get('id', "")

        # Get to action info
        to_action_data = action_info.get(to_action_type, {})
        to_action_id = to_action_data.get('id', "")
        to_port_data = to_action_data.get('ports', {}).get(to_port_usage, {})
        to_port_id = to_port_data.get('id', "")

        functional_exchanges.append({
            "exchange_id": conn_id,
            "exchange_name": conn_name,
            "from_action_id": from_action_id,
            "from_action_name": from_action_type,
            "from_port_id": from_port_id,
            "from_port_name": from_port_usage,  # Full name for ID lookup (e.g., FIP1_24b3)
            "to_action_id": to_action_id,
            "to_action_name": to_action_type,
            "to_port_id": to_port_id,
            "to_port_name": to_port_usage  # Full name for ID lookup (e.g., FOP1_ac6c)
        })

    return functional_exchanges

def create_workbook() -> openpyxl.Workbook:
    """Create a new workbook with headers."""
    wb = openpyxl.Workbook()

    if wb.active.title == "Sheet":
        wb.remove(wb.active)

    ws = wb.create_sheet("Functional Exchanges")

    headers = [
        "Function From ID", "Function From Name", "Function From Port ID", "Function From Port Name",
        "Functional Exchange ID", "Functional Exchange Name",
        "Function To ID", "Function To Name", "Function To Port ID", "Function To Port Name"
    ]

    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)

    for cell in ws[1]:
        cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

    return wb

def export_to_excel(functional_exchanges: List[Dict], output_path: pathlib.Path) -> None:
    """Export functional exchanges to an Excel file with formatted port names."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = create_workbook()
    ws = wb["Functional Exchanges"]

    for exchange in functional_exchanges:
        new_row = ws.max_row + 1

        # Format port names for export (remove ID suffix and add spaces)
        formatted_from_port_name = format_port_name_for_export(exchange["from_port_name"])
        formatted_to_port_name = format_port_name_for_export(exchange["to_port_name"])

        # From function and port
        ws.cell(row=new_row, column=1, value=exchange["from_action_id"])
        ws.cell(row=new_row, column=2, value=from_pascal_case(exchange["from_action_name"]))
        ws.cell(row=new_row, column=3, value=exchange["from_port_id"])
        ws.cell(row=new_row, column=4, value=formatted_from_port_name)  # Formatted port name

        # Functional exchange info
        ws.cell(row=new_row, column=5, value=exchange["exchange_id"])
        ws.cell(row=new_row, column=6, value=from_pascal_case(exchange["exchange_name"]))

        # To function and port
        ws.cell(row=new_row, column=7, value=exchange["to_action_id"])
        ws.cell(row=new_row, column=8, value=from_pascal_case(exchange["to_action_name"]))
        ws.cell(row=new_row, column=9, value=exchange["to_port_id"])
        ws.cell(row=new_row, column=10, value=formatted_to_port_name)  # Formatted port name

    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2) * 1.2
        ws.column_dimensions[column_letter].width = adjusted_width

    wb.save(output_path)
    print(f"Exported {len(functional_exchanges)} functional exchanges to: {output_path}")

def main() -> None:
    _DVS_DIR = pathlib.Path(__file__).parent.parent
    DEFAULT_INPUT = _DVS_DIR / "Functions_generated.sysml"
    DEFAULT_OUTPUT = _DVS_DIR / "results" / "DVS_Functional_Exchanges_export.xlsx"

    args = sys.argv[1:]
    if args and args[0] in ("-h", "--help"):
        print(__doc__)
        sys.exit(0)

    input_path = pathlib.Path(args[0]) if args else DEFAULT_INPUT
    if not input_path.exists():
        print(f"Error: file not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    output_path = DEFAULT_OUTPUT if len(args) < 2 else pathlib.Path(args[1])

    print(f"Reading: {input_path}")
    functional_exchanges = parse_sysml_file(input_path)

    if not functional_exchanges:
        print("Warning: No functional exchanges found in the file!")
    else:
        print(f"Found {len(functional_exchanges)} functional exchanges in the file")
        # Print debug info for the first few exchanges
        for i, exchange in enumerate(functional_exchanges[:3], 1):
            print(f"\nSample exchange {i}:")
            print(f"  From: {exchange['from_action_name']}.{format_port_name_for_export(exchange['from_port_name'])} (ID: {exchange['from_action_id']}, Port ID: {exchange['from_port_id']})")
            print(f"  To:   {exchange['to_action_name']}.{format_port_name_for_export(exchange['to_port_name'])} (ID: {exchange['to_action_id']}, Port ID: {exchange['to_port_id']})")
            print(f"  Exchange: {from_pascal_case(exchange['exchange_name'])} (ID: {exchange['exchange_id']})")

    export_to_excel(functional_exchanges, output_path)

if __name__ == "__main__":
    main()